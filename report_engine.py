# report_engine.py
import io
import re
import datetime as _dt
import pandas as pd
from openpyxl import Workbook

# ---------------- Helpers ----------------

def _read_csv(file_storage_or_path):
    """Accepts a Flask FileStorage or a file path. Returns a pandas DataFrame."""
    if hasattr(file_storage_or_path, "stream"):  # Flask FileStorage
        file_storage_or_path.stream.seek(0)
        return pd.read_csv(file_storage_or_path.stream)
    return pd.read_csv(file_storage_or_path)

def _sanitize_sheet_name(name: str) -> str:
    # Excel sheet name rules
    bad = r'[\\/*?:\[\]]'
    name = re.sub(bad, "_", str(name or "Sheet")).strip()
    return name[:31] or "Sheet"

def _write_df_to_sheet(wb: Workbook, sheet_name: str, df: pd.DataFrame):
    ws = wb.create_sheet(_sanitize_sheet_name(sheet_name))
    ws.append([str(c) for c in df.columns.tolist()])
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))
    return ws

def _ensure_report_sheet(wb: Workbook):
    # Replace default "Sheet" with "Report"
    if wb.worksheets:
        wb.remove(wb.worksheets[0])
    return wb.create_sheet("Report")

def _put_value(ws, cell_ref: str, value, number_format: str | None = None):
    cell = str(cell_ref).strip().upper()
    if not cell:
        return
    ws[cell] = value
    if number_format:
        ws[cell].number_format = number_format

def _put_formula(ws, cell_ref: str, formula_text: str):
    """Writes a formula (adds leading '=' if missing)."""
    cell = str(cell_ref).strip().upper()
    if not cell or not formula_text:
        return
    f = str(formula_text).strip()
    if not f.startswith("="):
        f = "=" + f
    ws[cell] = f

def _combine_gen_lims(gen_lims_csv_list):
    """
    Build two helper sheets from all Gen_LIMs_* CSVs:
      - Combined_CFM: columns [C, F, M] -> ["ID", "Analyte", "Result"]
      - Combined_CFW: columns [C, F, W] -> ["ID", "Analyte", "ValueW"]
    Returns (df_cfm, df_cfw).
    """
    frames_cfm, frames_cfw = [], []
    for item in gen_lims_csv_list or []:
        df = _read_csv(item)

        def _safe_pick_cols(df_, idxs):
            cols = []
            for idx in idxs:
                if idx < df_.shape[1]:
                    cols.append(df_.iloc[:, idx])
                else:
                    cols.append(pd.Series([""] * len(df_)))
            return pd.DataFrame({i: c for i, c in zip(idxs, cols)})

        # C(2), F(5), M(12); C(2), F(5), W(22)
        if df.shape[1] > 2:  # ensure col C exists
            cfm = _safe_pick_cols(df, [2, 5, 12])
            cfm.columns = ["ID", "Analyte", "Result"]
            frames_cfm.append(cfm)

            cfw = _safe_pick_cols(df, [2, 5, 22])
            cfw.columns = ["ID", "Analyte", "ValueW"]
            frames_cfw.append(cfw)

    df_cfm = pd.concat(frames_cfm, ignore_index=True) if frames_cfm else pd.DataFrame(columns=["ID","Analyte","Result"])
    df_cfw = pd.concat(frames_cfw, ignore_index=True) if frames_cfw else pd.DataFrame(columns=["ID","Analyte","ValueW"])
    return df_cfm, df_cfw

def _is_today_formula(text: str) -> bool:
    t = (text or "").strip().lower()
    return t in ("today()", "today", "=today()", "=today")

def _excel_equivalent(cell_ref: str, original_formula: str) -> str | None:
    """
    For a few known Google-only patterns, return an Excel/Sheets-compatible formula
    that uses our helper sheets (Combined_CFM / Combined_CFW).
    Extend this mapping as needed.
    """
    cell = (cell_ref or "").strip().upper()
    ftxt = (original_formula or "").strip()

    uses_google_only = ("GETSHEETNAMES" in ftxt) or ("QUERY(" in ftxt) or ("INDIRECT(" in ftxt)
    if not uses_google_only:
        return None

    # A28: First matching Analyte (BPS or PFAS) for current E17 from combined CFM
    if cell == "A28":
        return (
            "=IFERROR("
            "INDEX(FILTER(Combined_CFM!B:B,"
            "(LEFT(Combined_CFM!A:A,LEN($E$17))=$E$17)"
            "*((Combined_CFM!B:B=\"Bisphenol S\")+(Combined_CFM!B:B=\"PFAS\"))),1),"
            "\"Not Found\")"
        )

    # G28: First matching ValueW for current E17 from combined CFW
    if cell == "G28":
        return (
            "=IFERROR("
            "INDEX(FILTER(Combined_CFW!C:C,"
            "(LEFT(Combined_CFW!A:A,LEN($E$17))=$E$17)"
            "*((Combined_CFW!B:B=\"Bisphenol S\")+(Combined_CFW!B:B=\"PFAS\"))),1),"
            "\"Not Found\")"
        )

    # Add more explicit rewrites here if you add new Google-only constructs later.
    return None

# ---------------- Public API ----------------

def build_report_workbook(
    product_or_id_value: str,
    formulas_csv,                  # FileStorage or path; Column B = cell, Column C = formula
    total_products_csv,            # FileStorage or path -> sheet "TotalProducts"
    data_consolidator_csv,         # FileStorage or path -> sheet "Data_Consolidator"
    gen_lims_csv_list=None,        # list[FileStorage or path]; each -> own sheet; also used to build Combined_* helpers
    static_today: bool = True,     # if True, replace TODAY() with fixed date value
) -> io.BytesIO:
    """
    Returns an in-memory .xlsx with:
      - Sheet 'Report' with your formulas (with Google-only pieces rewritten or pre-filled when possible).
      - Sheet 'TotalProducts' from CSV.
      - Sheet 'Data_Consolidator' from CSV.
      - One sheet per uploaded Gen_LIMs_* CSV.
      - Helper sheets 'Combined_CFM' and 'Combined_CFW' (for Excel-friendly replacements).
      - Cell E17 pre-filled with product_or_id_value.
    """
    gen_lims_csv_list = gen_lims_csv_list or []

    # Load dataframes
    tp_df = _read_csv(total_products_csv)
    dc_df = _read_csv(data_consolidator_csv)
    f_df  = _read_csv(formulas_csv)

    if f_df.shape[1] < 3:
        raise ValueError("Formulas.csv must have at least 3 columns (B = target cell, C = formula text).")

    # Workbook
    wb = Workbook()
    report_ws = _ensure_report_sheet(wb)

    # Data sheets first
    _write_df_to_sheet(wb, "TotalProducts", tp_df)
    _write_df_to_sheet(wb, "Data_Consolidator", dc_df)

    # Gen_LIMs_* individual sheets
    for item in gen_lims_csv_list:
        name = getattr(item, "filename", None) or "Gen_LIMs_Data"
        name = name.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
        name = name.rsplit(".", 1)[0]
        _write_df_to_sheet(wb, name, _read_csv(item))

    # Helper sheets for Excel-friendly replacements
    df_cfm, df_cfw = _combine_gen_lims(gen_lims_csv_list)
    _write_df_to_sheet(wb, "Combined_CFM", df_cfm)
    _write_df_to_sheet(wb, "Combined_CFW", df_cfw)

    # Pre-fill dropdown (E17) with product/id selection
    if product_or_id_value:
        _put_value(report_ws, "E17", str(product_or_id_value))

    # Place formulas / values from Formulas.csv:
    # Column B -> cell, Column C -> formula or token (e.g., today(), dropdown)
    for _, row in f_df.iterrows():
        cell_ref  = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ""
        raw_text  = str(row.iloc[2]).strip() if pd.notna(row.iloc[2]) else ""

        if not cell_ref:
            continue

        # If it's a "dropdown" marker, skip (E17 already set)
        if raw_text.lower() in ("dropdown", "(dropdown)", "select"):
            continue

        # If it's TODAY(), optionally set a static date value
        if static_today and _is_today_formula(raw_text):
            _put_value(report_ws, cell_ref, _dt.date.today(), number_format="m/d/yyyy")
            continue

        # If the text parses as a date and doesn't start with '=', treat as a constant date value
        try:
            parsed = pd.to_datetime(raw_text, errors="raise")
            if not raw_text.startswith("="):
                _put_value(report_ws, cell_ref, parsed.date(), number_format="m/d/yyyy")
                continue
        except Exception:
            pass

        # Try to swap Google-only formulas for Excel-friendly equivalents
        replacement = _excel_equivalent(cell_ref, raw_text)
        if replacement:
            _put_formula(report_ws, cell_ref, replacement)
            continue

        # Otherwise, write the original formula as-is (Excel 365 supports FILTER/LET/SEARCH/etc.)
        _put_formula(report_ws, cell_ref, raw_text)

    # Optional: freeze a header region
    report_ws.freeze_panes = "A2"

    # Save to bytes
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
